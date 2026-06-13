"""Tests for Preview Service."""
from __future__ import annotations

import json
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch
from uuid import uuid4

import pytest
from fastapi import status

# Valid UUID for test file IDs
_TEST_FILE_ID = str(uuid4())


# ── Health & root ────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_health_check(client):
    """Health endpoint returns healthy when file-service is reachable."""
    mock_response = MagicMock()
    mock_response.status_code = 200

    with patch("src.main._http_client") as mock_client:
        mock_client.get = AsyncMock(return_value=mock_response)
        mock_client.is_closed = False

        resp = await client.get("/health")

    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "healthy"
    assert data["service"] == "preview"
    assert data["file_service"] == "healthy"


@pytest.mark.asyncio
async def test_health_check_degraded(client):
    """Health endpoint returns degraded when file-service is unreachable."""
    with patch("src.main._http_client") as mock_client:
        mock_client.get = AsyncMock(side_effect=Exception("connection refused"))
        mock_client.is_closed = False

        resp = await client.get("/health")

    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "degraded"
    assert data["file_service"] == "unreachable"


@pytest.mark.asyncio
async def test_root(client):
    """Service info endpoint returns metadata."""
    resp = await client.get("/api/preview/")
    assert resp.status_code == 200
    data = resp.json()
    assert data["version"] == "1.0.0"
    assert "txt" in data["supported_text_previews"]
    assert "docx" in data["supported_text_previews"]
    assert "xlsx" in data["supported_text_previews"]


# ── Auth ─────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_preview_requires_auth(client):
    """Preview endpoint returns 401 without Authorization header."""
    resp = await client.get(f"/api/preview/{_TEST_FILE_ID}")
    assert resp.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio
async def test_thumbnail_requires_auth(client):
    """501 stubs also require auth."""
    resp = await client.get(f"/api/preview/{_TEST_FILE_ID}/thumbnail")
    assert resp.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio
async def test_generate_requires_auth(client):
    resp = await client.post(f"/api/preview/{_TEST_FILE_ID}/generate")
    assert resp.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio
async def test_delete_requires_auth(client):
    resp = await client.delete(f"/api/preview/{_TEST_FILE_ID}")
    assert resp.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio
async def test_invalid_file_id_returns_400(client):
    """Non-UUID file_id is rejected with 400."""
    resp = await client.get(
        "/api/preview/not-a-uuid",
        headers={"Authorization": "Bearer test-token"},
    )
    assert resp.status_code == status.HTTP_400_BAD_REQUEST


# ── Mock helpers ─────────────────────────────────────────────────


def _make_file_response(content: bytes, mime_type: str, status_code: int = 200):
    """Build a mock httpx.Response-like object."""
    resp = MagicMock()
    resp.status_code = status_code
    resp.content = content
    resp.text = content.decode("utf-8", errors="replace")
    resp.headers = {"content-type": mime_type}
    return resp


def _patch_fetch(content: bytes, mime_type: str):
    """Context manager that patches _fetch_file_bytes."""
    return patch(
        "src.main._fetch_file_bytes",
        new_callable=AsyncMock,
        return_value=(content, mime_type),
    )


# ── Text preview ─────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_preview_plain_text(client):
    text = b"Hello, world!"
    with _patch_fetch(text, "text/plain"):
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 200
    data = resp.json()
    assert data["kind"] == "text"
    assert data["content"] == "Hello, world!"
    assert data["truncated"] is False


@pytest.mark.asyncio
async def test_preview_csv(client):
    csv_content = b"name,age\nAlice,30\nBob,25"
    with _patch_fetch(csv_content, "text/csv"):
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 200
    assert "Alice" in resp.json()["content"]


@pytest.mark.asyncio
async def test_preview_json_pretty_printed(client):
    raw = json.dumps({"key": "value", "nested": {"a": 1}}).encode()
    with _patch_fetch(raw, "application/json"):
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 200
    data = resp.json()
    assert "key" in data["content"]
    assert "\n" in data["content"]  # pretty-printed


# ── DOCX preview ─────────────────────────────────────────────────


def _make_docx_bytes(text: str) -> bytes:
    """Create a minimal DOCX in memory."""
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_preview_docx(client):
    docx_bytes = _make_docx_bytes("Test paragraph content")
    with _patch_fetch(docx_bytes, DOCX_MIME_TYPE):
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 200
    data = resp.json()
    assert "Test paragraph content" in data["content"]


DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── XLSX preview ─────────────────────────────────────────────────


def _make_xlsx_bytes() -> bytes:
    """Create a minimal XLSX in memory."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Value"])
    ws.append(["Alice", 100])
    ws.append(["Bob", 200])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_preview_xlsx(client):
    xlsx_bytes = _make_xlsx_bytes()
    with _patch_fetch(xlsx_bytes, XLSX_MIME_TYPE):
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 200
    data = resp.json()
    assert "Sheet1" in data["content"]
    assert "Alice" in data["content"]


# ── Unsupported type ─────────────────────────────────────────────


@pytest.mark.asyncio
async def test_preview_unsupported_type(client):
    with _patch_fetch(b"binary data", "application/octet-stream"):
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == status.HTTP_415_UNSUPPORTED_MEDIA_TYPE


# ── Truncation ───────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_preview_text_truncation(client):
    large_text = ("A" * 50000).encode()
    with _patch_fetch(large_text, "text/plain"):
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 200
    data = resp.json()
    assert data["truncated"] is True
    assert len(data["content"]) == 40000


# ── File-service errors ──────────────────────────────────────────


@pytest.mark.asyncio
async def test_preview_file_not_found(client):
    """Propagates 404 from file-service."""
    with patch("src.main._fetch_file_bytes", new_callable=AsyncMock) as mock_fetch:
        from fastapi import HTTPException

        mock_fetch.side_effect = HTTPException(
            status_code=404, detail="File not found"
        )
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_preview_file_forbidden(client):
    """Propagates 403 from file-service."""
    with patch("src.main._fetch_file_bytes", new_callable=AsyncMock) as mock_fetch:
        from fastapi import HTTPException

        mock_fetch.side_effect = HTTPException(
            status_code=403, detail="Access denied"
        )
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 403


# ── X-API-Key ────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_preview_forwards_api_key(client):
    """When SERVICE_API_KEY is set, it's sent as X-API-Key."""
    with patch("src.main.settings") as mock_settings:
        mock_settings.file_service_url = "http://file:8000"
        mock_settings.service_api_key = "test-secret-key"
        mock_settings.preview_max_size = 10485760

        with patch("src.main._http_client") as mock_client:
            mock_resp = MagicMock()
            mock_resp.status_code = 200
            mock_resp.headers = {"content-type": "text/plain"}

            # Mock streaming response
            async def aiter_bytes(chunk_size):
                yield b"hello"
            mock_resp.aiter_bytes = aiter_bytes

            mock_client.get = AsyncMock(return_value=mock_resp)
            mock_client.is_closed = False

            await client.get(
                f"/api/preview/{_TEST_FILE_ID}",
                headers={"Authorization": "Bearer test-token"},
            )

            # Verify X-API-Key was sent
            call_args = mock_client.get.call_args
            assert call_args[1]["headers"]["X-API-Key"] == "test-secret-key"


# ── Rate limiting ─────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_rate_limit_blocks_after_max_requests(client):
    """After _RATE_LIMIT_MAX_REQUESTS, preview returns 429."""
    import src.main as main_mod

    # Reset rate limit store for this key
    test_key = "Bearer test-token"[:16]
    main_mod._rate_limit_store[test_key] = []

    with _patch_fetch(b"content", "text/plain"):
        # Exhaust the rate limit
        for _ in range(main_mod._RATE_LIMIT_MAX_REQUESTS):
            resp = await client.get(
                f"/api/preview/{_TEST_FILE_ID}",
                headers={"Authorization": "Bearer test-token"},
            )
            assert resp.status_code == 200

        # Next request should be rate-limited
        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == status.HTTP_429_TOO_MANY_REQUESTS
    assert "Rate limit exceeded" in resp.json()["detail"]

    # Clean up
    main_mod._rate_limit_store.pop(test_key, None)


# ── Error sanitization ────────────────────────────────────────────


@pytest.mark.asyncio
async def test_upstream_5xx_error_is_sanitized(client):
    """5xx errors from file-service should return generic message."""
    mock_resp = MagicMock()
    mock_resp.status_code = 500
    mock_resp.text = "Internal DB connection pool exhausted"
    mock_resp.headers = {"content-type": "text/plain"}

    with patch("src.main._http_client") as mock_client:
        mock_client.get = AsyncMock(return_value=mock_resp)
        mock_client.is_closed = False

        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 500
    # The detail should NOT leak the internal error message
    assert "DB connection" not in resp.json()["detail"]
    assert resp.json()["detail"] == "File service error"


@pytest.mark.asyncio
async def test_upstream_4xx_error_preserved(client):
    """4xx errors from file-service should preserve the detail."""
    mock_resp = MagicMock()
    mock_resp.status_code = 403
    mock_resp.text = "Access denied"
    mock_resp.headers = {"content-type": "text/plain"}

    with patch("src.main._http_client") as mock_client:
        mock_client.get = AsyncMock(return_value=mock_resp)
        mock_client.is_closed = False

        resp = await client.get(
            f"/api/preview/{_TEST_FILE_ID}",
            headers={"Authorization": "Bearer test-token"},
        )

    assert resp.status_code == 403
    assert resp.json()["detail"] == "Access denied"
